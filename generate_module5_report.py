from docx import Document
from docx.shared import Pt

doc = Document()

def heading(text, level):
    doc.add_heading(text, level=level)

def paragraph(text='', bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p

def blank_placeholder(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    return p

# Title
heading('RAPPORT COMPLET - MODULE 5', 0)
paragraph('Communium - Système de gestion d\'événements', bold=False)
paragraph()
paragraph('Ce document contient un état complet du Module 5, avec des sections réservées pour insérer des captures d\'écran frontend et un schéma global expliquant les démarches de la solution.', bold=False)
paragraph()

# Table des matières
heading('1. Table des matières', 1)
for item in [
    '1. Table des matières',
    '2. Vue d\'ensemble',
    '3. Architecture de la solution',
    '4. Schéma global',
    '5. Fonctionnalités implémentées',
    '6. Interfaces utilisateur',
    '7. Processus d\'inscription',
    '8. Gestion des billets',
    '9. Détails techniques',
    '10. Base de données',
    '11. Résultats et validations',
    '12. Conclusion',
    '13. Appendices'
]:
    paragraph(item)

# Overview
heading('2. Vue d\'ensemble', 1)
paragraph('Objectif :')
paragraph('Le Module 5 est un système complet de gestion d\'événements permettant aux organisateurs de créer, gérer et contrôler des événements professionnels.', bold=False)

paragraph('Périmètre fonctionnel :')
for item in [
    'Listing des événements publiés',
    'Création d\'événements',
    'Inscription gratuite avec génération de billet',
    'Inscription payante avec intention de paiement',
    'Génération de PDF avec QR code',
    'Envoi de billet par email',
    'Gestion de la capacité et listes d\'attente',
    'Check-in des participants'
]:
    paragraph(f'• {item}')

# Architecture
heading('3. Architecture de la solution', 1)
paragraph('Stack technologique :')
for item in [
    'Frontend : Next.js (App Router)',
    'Backend : Node.js + Express',
    'Base de données : PostgreSQL',
    'Email : Nodemailer',
    'QR Code : qrcode',
    'PDF : pdfkit'
]:
    paragraph(f'• {item}')
paragraph()
paragraph('Ports :')
paragraph('• Frontend : http://localhost:3001')
paragraph('• Backend : http://localhost:5000')
paragraph('• PostgreSQL : postgresql://localhost:5432/communium')

# Schema global placeholder
heading('4. Schéma global', 1)
paragraph('Cette section est réservée pour le schéma global expliquant les démarches de la solution.', bold=False)
blank_placeholder('[PLACEHOLDER - INSÉRER LE SCHÉMA GLOBAL ICI]')
paragraph('Le schéma doit montrer les interactions entre :')
for item in [
    'Frontend',
    'Backend API',
    'Base de données PostgreSQL',
    'Service email',
    'Fournisseurs de paiement',
    'Utilisateur final'
]:
    paragraph(f'• {item}')
paragraph('Utiliser ce bloc pour expliquer les flux de création d\'événements, d\'inscription, de génération de billets et de check-in.', bold=False)

# Functionalities
heading('5. Fonctionnalités implémentées', 1)
paragraph('5.1 Gestion des événements :')
paragraph('• Listing', bold=False)
paragraph('• Détail', bold=False)
paragraph('• Création', bold=False)
paragraph('• Mise à jour', bold=False)

paragraph('5.2 Inscription aux événements :')
paragraph('• Flux gratuit', bold=False)
paragraph('• Flux payant', bold=False)
paragraph('• Gestion des files d\'attente', bold=False)

paragraph('5.3 Gestion des billets :')
paragraph('• Génération PDF', bold=False)
paragraph('• Génération QR code', bold=False)
paragraph('• Envoi par email', bold=False)
paragraph('• Check-in', bold=False)

# Interfaces utilisateurs
heading('6. Interfaces utilisateur', 1)
paragraph('6.1 Page d\'accueil Module 5', bold=False)
blank_placeholder('[PLACEHOLDER - INSÉRER CAPTURE D\'ÉCRAN: LISTING DES ÉVÉNEMENTS]')
paragraph('6.2 Détail d\'un événement', bold=False)
blank_placeholder('[PLACEHOLDER - INSÉRER CAPTURE D\'ÉCRAN: DÉTAIL D\'UN ÉVÉNEMENT GRATUIT]')
blank_placeholder('[PLACEHOLDER - INSÉRER CAPTURE D\'ÉCRAN: DÉTAIL D\'UN ÉVÉNEMENT PAYANT]')
paragraph('6.3 Confirmation d\'inscription', bold=False)
blank_placeholder('[PLACEHOLDER - INSÉRER CAPTURE D\'ÉCRAN: TICKET GÉNÉRÉ]')

# Processus inscription
heading('7. Processus d\'inscription', 1)
paragraph('7.1 Flux d\'inscription gratuite', bold=False)
paragraph('Décrire le parcours utilisateur et les appels API pour les billets gratuits.', bold=False)
paragraph('7.2 Flux d\'inscription payante', bold=False)
paragraph('Décrire le parcours utilisateur et le workflow de paiement.', bold=False)
paragraph('7.3 Gestion de la file d\'attente', bold=False)
paragraph('Décrire le processus de mise en attente et de promotion.', bold=False)

# Gestion billets
heading('8. Gestion des billets', 1)
paragraph('8.1 Composition du billet PDF', bold=False)
paragraph('• Titre de l\'événement', bold=False)
paragraph('• Nom et email du participant', bold=False)
paragraph('• Code de billet unique', bold=False)
paragraph('• Date et heure de l\'événement', bold=False)
paragraph('• QR code', bold=False)
paragraph('8.2 Stockage des billets', bold=False)
paragraph('• Table : module5.event_tickets', bold=False)

# Détails techniques
heading('9. Détails techniques', 1)
paragraph('9.1 Endpoints API', bold=False)
paragraph('• GET /api/module5/events', bold=False)
paragraph('• GET /api/module5/events/:eventId', bold=False)
paragraph('• POST /api/module5/events', bold=False)
paragraph('• POST /api/module5/events/:eventId/register', bold=False)
paragraph('• GET /api/module5/tickets/:ticketId/pdf', bold=False)
paragraph('• POST /api/module5/tickets/:ticketId/email', bold=False)
paragraph('• POST /api/module5/tickets/checkin', bold=False)

# Base de données
heading('10. Base de données', 1)
paragraph('10.1 Schéma module5', bold=False)
paragraph('Inclure les tables principales et leurs relations.', bold=False)
paragraph('10.2 Données de test', bold=False)
paragraph('Lister les événements de test et les informations clés.', bold=False)

# Résultats
heading('11. Résultats et validations', 1)
paragraph('Inclure les tests réalisés, les validations et l\'état actuel du système.', bold=False)
paragraph('• Listing des événements', bold=False)
paragraph('• Détail d\'un événement', bold=False)
paragraph('• Inscription gratuite', bold=False)
paragraph('• Génération PDF + QR', bold=False)
paragraph('• Envoi d\'email', bold=False)
paragraph('• CORS Frontend-Backend', bold=False)

# Conclusion
heading('12. Conclusion', 1)
paragraph('Récapitulatif des réalisations et prochaines étapes.', bold=False)

# Appendices
heading('13. Appendices', 1)
paragraph('Variables d\'environnement', bold=False)
paragraph('• Backend: .env', bold=False)
paragraph('• Frontend: .env.local', bold=False)
paragraph('Commandes utiles', bold=False)
paragraph('• Démarrer le backend', bold=False)
paragraph('• Démarrer le frontend', bold=False)
paragraph('• Tester les endpoints', bold=False)
paragraph('• Vérifier la base PostgreSQL', bold=False)

# Save
path = 'MODULE5_RAPPORT_COMPLET.docx'
doc.save(path)
print(f'Generated {path}')
